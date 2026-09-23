"""Экспорт итогового протокола: DOCX (python-docx) и PDF (reportlab, кириллица)."""
from __future__ import annotations

import io
import re
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.units import cm as RLcm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ACCENT = colors.HexColor("#1B5B90")
DARK = colors.HexColor("#222831")

_FONT_DIRS = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "/usr/share/fonts/truetype/freefont/FreeSans.ttf",
]
_REGISTERED: dict[str, str] = {}


def _find_font() -> list[str]:
    import glob
    for pat in ("/usr/share/fonts/**/DejaVuSans.ttf",
                "/usr/share/fonts/**/DejaVuSans-Bold.ttf"):
        hits = sorted(glob.glob(pat, recursive=True))
        if hits:
            return hits[:2]
    return _FONT_DIRS[:2]


def _register() -> tuple[str, str]:
    if _REGISTERED:
        return _REGISTERED["normal"], _REGISTERED["bold"]
    found = _find_font()
    normal = found[0] if found else None
    bold = found[1] if len(found) > 1 else normal
    if normal:
        pdfmetrics.registerFont(TTFont("Khattama", normal))
        if bold:
            pdfmetrics.registerFont(TTFont("Khattama-Bold", bold))
    _REGISTERED["normal"], _REGISTERED["bold"] = "Khattama", "Khattama-Bold"
    return "Khattama", "Khattama-Bold"


def _ts(sec: float) -> str:
    m, s = divmod(int(sec), 60)
    return f"{m:02d}:{s:02d}"


def _speaker_color_hash(name: str) -> str:
    palette = ["#1B5B90", "#2E7D32", "#C62828", "#6A1B9A", "#E65100",
               "#00838F", "#37474F", "#33691E", "#880E4F", "#5D4037"]
    h = sum(ord(c) for c in name)
    return palette[h % len(palette)]


def build_text_protocol(meeting: dict) -> str:
    lines: list[str] = []
    lines.append("ПРОТОКОЛ СОВЕЩАНИЯ (ИИ)")
    lines.append("=" * 60)
    lines.append(f"Файл: {meeting.get('filename', '—')}")
    lines.append(f"Обработано: {meeting.get('created_at', '—')}")
    lines.append(f"Источник поручений: {meeting.get('extract_mode', '—')}")
    lines.append("")

    lines.append("КРАТКОЕ САММАРИ")
    lines.append("-" * 60)
    for s in meeting.setdefault("summary", []):
        lines.append(f"• {s}")
    lines.append("")

    lines.append("ПОРУЧЕНИЯ")
    lines.append("-" * 60)
    a = meeting.get("assignments", [])
    if not a:
        lines.append("— поручения не обнаружены —")
    for i, p in enumerate(a, 1):
        deadline = p.get("deadline") or p.get("deadline_raw") or "не указан"
        owner = p.get("owner") or p.get("speaker_label") or "не указан"
        lines.append(f"{i}. [{p.get('priority', 'medium').upper()}] {p.get('task', '')}")
        lines.append(f"   Ответственный: {owner} | Срок: {deadline} | Говорящий: {p.get('speaker_label') or '—'}")
    lines.append("")

    lines.append("ТРАНСКРИПТ (по говорящим)")
    lines.append("-" * 60)
    for b in meeting.get("transcript", []):
        lines.append(f"[{_ts(b['start'])}–{_ts(b['end'])}] {b.get('speaker', '—')}: {b.get('text', '')}")
    return "\n".join(lines)


def build_docx(meeting: dict) -> io.BytesIO:
    buf = io.BytesIO()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    h = doc.add_heading("Протокол совещания (Хаттама AI)", level=0)
    h.runs[0].font.color.rgb = RGBColor(0x1B, 0x5B, 0x90)
    doc.add_paragraph(f"Файл: {meeting.get('filename', '—')}   ·   Обработано: {meeting.get('created_at', '—')}")
    doc.add_paragraph(f"Поручения выделены: {meeting.get('extract_mode', '—')} · Саммари: {meeting.get('summary_mode', '—')}")

    doc.add_heading("Краткое саммари", level=1)
    for s in meeting.get("summary", []):
        doc.add_paragraph(s, style="List Bullet")

    doc.add_heading("Поручения", level=1)
    a = meeting.get("assignments", [])
    if a:
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = table.rows[0].cells
        for i, name in enumerate(["№", "Поручение", "Ответственный", "Срок", "Приоритет"]):
            hdr[i].text = name
        for i, p in enumerate(a, 1):
            cells = table.add_row().cells
            cells[0].text = str(i)
            cells[1].text = p.get("task", "")
            cells[2].text = p.get("owner") or p.get("speaker_label") or "не указан"
            cells[3].text = p.get("deadline") or p.get("deadline_raw") or "не указан"
            cells[4].text = p.get("priority", "medium")
    else:
        doc.add_paragraph("— поручения не обнаружены —")

    doc.add_heading("Транскрипт по говорящим", level=1)
    for b in meeting.get("transcript", []):
        p = doc.add_paragraph()
        run = p.add_run(f"[{_ts(b['start'])}–{_ts(b['end'])}] {b.get('speaker', '')}: ")
        run.bold = True
        p.add_run(b.get("text", ""))

    doc.save(buf)
    buf.seek(0)
    return buf


def build_pdf(meeting: dict) -> io.BytesIO:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=RLcm * 2, rightMargin=RLcm * 2,
                            topMargin=RLcm * 2, bottomMargin=RLcm * 2)
    styles = getSampleStyleSheet()
    normal_font, bold_font = _register()

    title = Paragraph("Протокол совещания — Хаттама AI",
                      ParagraphStyle("titleK", fontName=bold_font, fontSize=18,
                                     textColor=DARK, spaceAfter=4))
    meta = Paragraph(
        f"Файл: {meeting.get('filename', '—')}<br/>Обработано: {meeting.get('created_at', '—')}",
        ParagraphStyle("metaK", fontName=normal_font, fontSize=9, textColor=colors.grey, spaceAfter=12))

    story = [title, meta]

    story.append(Paragraph("Краткое саммари",
                           ParagraphStyle("hK1", fontName=bold_font, fontSize=13,
                                          textColor=ACCENT, spaceBefore=10, spaceAfter=4)))
    for s in meeting.get("summary", []):
        story.append(Paragraph(f"• {s}", ParagraphStyle("sumK", fontName=normal_font,
                                                        fontSize=10.5, spaceAfter=2)))

    story.append(Paragraph("Поручения",
                           ParagraphStyle("hK2", fontName=bold_font, fontSize=13,
                                          textColor=ACCENT, spaceBefore=10, spaceAfter=4)))
    a = meeting.get("assignments", [])
    if a:
        data = [["№", "Поручение", "Ответственный", "Срок", "Приоритет"]]
        for i, p in enumerate(a, 1):
            data.append([str(i), p.get("task", ""),
                         p.get("owner") or p.get("speaker_label") or "не указан",
                         p.get("deadline") or p.get("deadline_raw") or "не указан",
                         p.get("priority", "medium")])
        widths = [RLcm, RLcm * 8, RLcm * 3.2, RLcm * 3.2, RLcm * 1.6]
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), ACCENT),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), bold_font),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("FONTNAME", (0, 1), (-1, -1), normal_font),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B0BEC5")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F7FA")]),
        ]))
        story.append(t)
    else:
        story.append(Paragraph("— поручения не обнаружены —",
                               ParagraphStyle("nK", fontName=normal_font, fontSize=10.5)))

    story.append(Paragraph("Транскрипт по говорящим",
                           ParagraphStyle("hK3", fontName=bold_font, fontSize=13,
                                          textColor=ACCENT, spaceBefore=12, spaceAfter=4)))
    for b in meeting.get("transcript", []):
        color = _speaker_color_hash(b.get("speaker", "?"))
        p = Paragraph(
            f"<font color='{color}'><b>[{_ts(b['start'])}–{_ts(b['end'])}] "
            f"{b.get('speaker', '')}:</b></font> {b.get('text', '')}",
            ParagraphStyle("dK", fontName=normal_font, fontSize=9.5, leading=12,
                           spaceAfter=3))
        story.append(p)

    doc.build(story)
    buf.seek(0)
    return buf